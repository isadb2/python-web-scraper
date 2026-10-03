from __future__ import annotations

import csv
import logging
import re
import time
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Iterable
from urllib.parse import urljoin

import requests
from bs4 import BeautifulSoup
from requests.adapters import HTTPAdapter
from urllib3.util.retry import Retry

LOGGER = logging.getLogger(__name__)

BASE_URL = "https://books.toscrape.com/"
DEFAULT_DELAY = 0.8
DEFAULT_TIMEOUT = 20

RATING_MAP = {"One": 1, "Two": 2, "Three": 3, "Four": 4, "Five": 5}


@dataclass(slots=True)
class Book:
    title: str
    price_gbp: float
    availability: str
    rating: int
    upc: str = ""
    stock_count: int | None = None
    category: str = ""
    product_url: str = ""
    description: str = ""


def build_session() -> requests.Session:
    """Create a session with conservative retry behavior."""
    session = requests.Session()
    retry = Retry(
        total=3,
        connect=3,
        read=3,
        backoff_factor=0.8,
        status_forcelist=(429, 500, 502, 503, 504),
        allowed_methods=frozenset({"GET"}),
        respect_retry_after_header=True,
    )
    adapter = HTTPAdapter(max_retries=retry)
    session.mount("https://", adapter)
    session.mount("http://", adapter)
    session.headers.update(
        {
            "User-Agent": "AI-Freelance-Scraper/1.0 (portfolio demo; polite scraper)",
            "Accept-Language": "en-US,en;q=0.9",
        }
    )
    return session


def fetch_html(session: requests.Session, url: str, timeout: int = DEFAULT_TIMEOUT) -> str:
    """Fetch HTML and raise a useful exception on HTTP/network failure."""
    LOGGER.info("GET %s", url)
    response = session.get(url, timeout=timeout)
    response.raise_for_status()
    if "text/html" not in response.headers.get("Content-Type", ""):
        raise ValueError(f"Expected HTML response, got {response.headers.get('Content-Type')!r}")
    return response.text


def parse_price(text: str) -> float:
    match = re.search(r"([0-9]+(?:\.[0-9]+)?)", text)
    if not match:
        raise ValueError(f"Could not parse price from {text!r}")
    return float(match.group(1))


def parse_rating(classes: Iterable[str]) -> int:
    for name in classes:
        if name in RATING_MAP:
            return RATING_MAP[name]
    raise ValueError(f"Could not parse rating from classes: {list(classes)!r}")


def parse_stock(text: str) -> tuple[str, int | None]:
    clean = " ".join(text.split())
    if "In stock" not in clean:
        return "Out of stock", 0
    match = re.search(r"(\d+)\s+available", clean, flags=re.I)
    return "In stock", int(match.group(1)) if match else None


def parse_listing_page(html: str, base_url: str) -> tuple[list[Book], str | None]:
    soup = BeautifulSoup(html, "html.parser")
    books: list[Book] = []

    for card in soup.select("article.product_pod"):
        link = card.select_one("h3 a")
        price = card.select_one(".price_color")
        availability = card.select_one(".availability")
        star = card.select_one(".star-rating")

        if not all((link, price, availability, star)):
            LOGGER.warning("Skipping malformed product card")
            continue

        title = (link.get("title") or link.get_text(" ", strip=True)).strip()
        product_url = urljoin(base_url, link.get("href", ""))
        rating = parse_rating(star.get("class", []))
        stock_status, stock_count = parse_stock(availability.get_text(" ", strip=True))

        books.append(
            Book(
                title=title,
                price_gbp=parse_price(price.get_text(" ", strip=True)),
                availability=stock_status,
                rating=rating,
                stock_count=stock_count,
                product_url=product_url,
            )
        )

    next_link = soup.select_one("li.next a")
    next_url = urljoin(base_url, next_link.get("href")) if next_link and next_link.get("href") else None
    return books, next_url


def parse_detail_page(book: Book, html: str) -> Book:
    soup = BeautifulSoup(html, "html.parser")
    table = {
        row.select_one("th").get_text(" ", strip=True): row.select_one("td").get_text(" ", strip=True)
        for row in soup.select("table.table tr")
        if row.select_one("th") and row.select_one("td")
    }

    description = ""
    heading = soup.find(id="product_description")
    if heading:
        paragraph = heading.find_next("p")
        if paragraph:
            description = paragraph.get_text(" ", strip=True)

    book.upc = table.get("UPC", "")
    book.category = table.get("Product Type", "")

    number_match = re.search(r"(\d+)\s+available", table.get("Availability", ""), flags=re.I)
    if number_match:
        book.stock_count = int(number_match.group(1))
        book.availability = "In stock"
    elif table.get("Availability"):
        book.availability = "Out of stock" if "out" in table["Availability"].lower() else book.availability

    book.description = description
    return book


def scrape_catalog(
    pages: int | None = 3,
    *,
    session: requests.Session | None = None,
    delay: float = DEFAULT_DELAY,
    timeout: int = DEFAULT_TIMEOUT,
    max_detail_pages: int | None = None,
) -> list[Book]:
    """Scrape catalog pages and enrich each listing with detail-page data.

    `pages=None` follows pagination until there is no next page.
    """
    owns_session = session is None
    session = session or build_session()
    results: list[Book] = []
    current_url = BASE_URL
    page_number = 0
    detail_count = 0

    try:
        while current_url and (pages is None or page_number < pages):
            page_number += 1
            LOGGER.info("Scraping listing page %s", page_number)
            html = fetch_html(session, current_url, timeout)
            page_books, next_url = parse_listing_page(html, current_url)
            LOGGER.info("Found %s products on page %s", len(page_books), page_number)

            for book in page_books:
                if max_detail_pages is not None and detail_count >= max_detail_pages:
                    results.append(book)
                    continue
                try:
                    time.sleep(delay)
                    detail_html = fetch_html(session, book.product_url, timeout)
                    parse_detail_page(book, detail_html)
                    detail_count += 1
                except requests.RequestException as exc:
                    LOGGER.warning("Detail page failed for %s: %s", book.product_url, exc)
                results.append(book)

            if next_url:
                time.sleep(delay)
            current_url = next_url
    finally:
        if owns_session:
            session.close()

    return results


def save_csv(books: list[Book], path: str | Path) -> Path:
    output = Path(path)
    output.parent.mkdir(parents=True, exist_ok=True)
    fieldnames = list(Book.__dataclass_fields__.keys())
    with output.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(asdict(book) for book in books)
    return output


def save_xlsx(books: list[Book], path: str | Path) -> Path:
    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter

    output = Path(path)
    output.parent.mkdir(parents=True, exist_ok=True)
    fieldnames = list(Book.__dataclass_fields__.keys())

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Books"
    sheet.append(fieldnames)
    for book in books:
        sheet.append([getattr(book, name) for name in fieldnames])

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for idx, name in enumerate(fieldnames, start=1):
        width = max(len(name), *(len(str(sheet.cell(r, idx).value or "")) for r in range(1, sheet.max_row + 1)))
        sheet.column_dimensions[get_column_letter(idx)].width = min(width + 2, 50)

    workbook.save(output)
    return output
